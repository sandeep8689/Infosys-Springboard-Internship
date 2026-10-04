import json
import re
from io import BytesIO

from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl import Workbook

from backend.database import get_consultation


def _safe_text(text: str) -> str:
    if not text:
        return ""
    text = str(text)
    text = text.replace("\u2013", "-").replace("\u2014", "-").replace("\u2019", "'")
    text = text.encode("latin-1", errors="replace").decode("latin-1")
    return re.sub(r"[^\x00-\xff]", "?", text)


def _parse_row(row: dict) -> dict:
    symptoms = json.loads(row["symptoms_json"])
    rec_raw = row["recommendations_json"]
    rec = json.loads(rec_raw) if isinstance(rec_raw, str) else rec_raw
    if not isinstance(rec, dict):
        rec = {}
    severity = row.get("severity_score")
    if severity is None:
        severity = rec.get("severity_score")
    return {
        "id": row["id"],
        "patient_name": row["patient_name"],
        "age": row["age"],
        "symptoms": symptoms,
        "predicted_disease": row["predicted_disease"],
        "confidence": float(row["confidence"]),
        "risk_level": row["risk_level"],
        "severity_score": float(severity) if severity is not None else None,
        "recommendations": rec,
        "created_at": row["created_at"],
    }


def _pdf_line(pdf: FPDF, text: str, h: float = 6, bold: bool = False) -> None:
    pdf.set_font("Helvetica", "B" if bold else "", 10 if not bold else 11)
    pdf.multi_cell(w=0, h=h, text=_safe_text(text), new_x=XPos.LMARGIN, new_y=YPos.NEXT)


def _pdf_bytes(pdf: FPDF) -> bytes:
    raw = pdf.output()
    if isinstance(raw, (bytes, bytearray)):
        data = bytes(raw)
    elif isinstance(raw, str):
        data = raw.encode("latin-1", errors="replace")
    else:
        raise ValueError("Unexpected PDF output type")
    if not data.startswith(b"%PDF"):
        raise ValueError("Invalid PDF document generated")
    return data


async def build_health_report_pdf(consultation_id: int) -> bytes:
    row = await get_consultation(consultation_id)
    if not row:
        raise ValueError("Consultation not found")
    data = _parse_row(row)

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.set_text_color(0, 128, 128)
    pdf.cell(0, 10, "MedAssist", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(60, 60, 60)
    pdf.cell(
        0,
        6,
        "AI Medical Symptom Analysis - Health Summary Report",
        new_x=XPos.LMARGIN,
        new_y=YPos.NEXT,
    )
    pdf.ln(4)
    pdf.set_text_color(0, 0, 0)

    _pdf_line(pdf, "Patient information", bold=True)
    _pdf_line(
        pdf,
        f"Name: {data['patient_name']}\n"
        f"Age: {data['age'] if data['age'] is not None else 'N/A'}\n"
        f"Consultation ID: {data['id']}\n"
        f"Generated (UTC): {data['created_at']}",
    )
    pdf.ln(2)

    _pdf_line(pdf, "Prediction, risk & severity", bold=True)
    sev = data["severity_score"]
    sev_txt = f"{sev:.1f} / 100" if sev is not None else "N/A"
    _pdf_line(
        pdf,
        f"Disease: {data['predicted_disease']}\n"
        f"Confidence: {data['confidence']:.1%}\n"
        f"Risk level: {data['risk_level']}\n"
        f"Severity score: {sev_txt}",
    )
    pdf.ln(2)

    _pdf_line(pdf, "Symptoms reported", bold=True)
    _pdf_line(pdf, ", ".join(data["symptoms"]) or "None listed")

    rec = data["recommendations"]
    for title, key in [
        ("Preventive care", "preventive_care"),
        ("Lifestyle advice", "lifestyle_advice"),
        ("Follow-up guidance", "follow_up_guidance"),
    ]:
        pdf.ln(2)
        _pdf_line(pdf, title, bold=True)
        items = rec.get(key) or []
        if not items:
            _pdf_line(pdf, "  - No items recorded.")
        for item in items:
            _pdf_line(pdf, f"  - {item}")

    if rec.get("llm_summary"):
        pdf.ln(2)
        _pdf_line(pdf, "Additional notes", bold=True)
        _pdf_line(pdf, rec["llm_summary"])

    pdf.ln(4)
    pdf.set_font("Helvetica", "I", 8)
    pdf.multi_cell(
        w=0,
        h=4,
        text=_safe_text(
            "Disclaimer: MedAssist is for educational use only. "
            "Not a substitute for professional medical advice."
        ),
        new_x=XPos.LMARGIN,
        new_y=YPos.NEXT,
    )

    return _pdf_bytes(pdf)


async def build_health_report_excel(consultation_id: int) -> bytes:
    row = await get_consultation(consultation_id)
    if not row:
        raise ValueError("Consultation not found")
    data = _parse_row(row)

    wb = Workbook()
    ws = wb.active
    ws.title = "MedAssist Report"
    ws.append(["MedAssist - Health Summary"])
    ws.append(["Field", "Value"])
    ws.append(["Consultation ID", data["id"]])
    ws.append(["Patient", data["patient_name"]])
    ws.append(["Age", data["age"]])
    ws.append(["Predicted disease", data["predicted_disease"]])
    ws.append(["Confidence", round(data["confidence"], 4)])
    ws.append(["Risk level", data["risk_level"]])
    ws.append(["Severity score (0-100)", data["severity_score"]])
    ws.append(["Symptoms", ", ".join(data["symptoms"])])
    ws.append([])
    rec = data["recommendations"]
    for section in ("preventive_care", "lifestyle_advice", "follow_up_guidance"):
        ws.append([section.replace("_", " ").title()])
        for item in rec.get(section, []) or []:
            ws.append(["", item])
        ws.append([])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
